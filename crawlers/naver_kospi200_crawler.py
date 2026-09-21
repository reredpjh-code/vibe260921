"""
PyQt6 GUI 기반 코스피 200 시세 크롤러

https://stock.naver.com/market/stock/kr 는 Next.js로 만들어진 클라이언트 렌더링(SPA)
페이지라 requests로 받은 HTML에는 시세 표가 들어있지 않다(BeautifulSoup로 파싱할
<table>이 없음). 예전에 표를 그대로 내려주던 finance.naver.com/sise/... 페이지들도
현재는 같은 신규 앱으로 통합되었거나(entryJongmok.naver는 HTTP 410) 더 이상 정적
HTML을 반환하지 않는다.

대신 그 화면이 실제로 데이터를 채울 때 호출하는, 네이버 증권이 공식적으로 서비스
중인 모바일 JSON API(m.stock.naver.com/api/...)를 그대로 사용한다. 이 API는
시가총액 내림차순으로 종목을 반환하는데, "코스피" 카테고리에는 일반 주식 외에
ETF/ETN도 함께 섞여 나온다. 각 항목의 stockEndType 값("stock"/"etf"/"etn")으로
ETF·ETN을 걸러내고, 일반 주식만 시가총액 상위 200개까지 모아 '코스피 200'으로
사용한다(코스피200 지수의 공식 편입 종목 명단과 100% 일치를 보장하지는 않지만,
네이버가 더 이상 공식 편입 종목 리스트 페이지를 제공하지 않는 상황에서 가장 가까운
근사치다).
"""
import csv
import sys

import requests
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtGui import QColor
from PyQt6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
}

MARKET_VALUE_URL = "https://m.stock.naver.com/api/stocks/marketValue/KOSPI"
DETAIL_URL_TEMPLATE = "https://stock.naver.com/domestic/stock/{code}"

TARGET_COUNT = 200
PAGE_SIZE = 100  # API가 허용하는 최대 페이지 크기

SIGN_BY_TEXT = {"상승": "+"}  # API가 하락일 때는 이미 '-'를 붙여서 내려준다


def fetch_kospi200(progress_callback=None):
    """코스피 '일반 주식'을 시가총액 상위 TARGET_COUNT개까지 모아서 반환한다.

    ETF/ETN은 stockEndType으로 걸러내므로, 목표 개수를 채울 때까지 필요한 만큼
    페이지를 계속 더 가져온다(ETF 비중이 큰 구간을 지나면 더 많은 페이지가 필요).
    """
    stocks = []
    page = 1
    max_pages = 30  # 안전장치: 무한 루프 방지
    while len(stocks) < TARGET_COUNT and page <= max_pages:
        res = requests.get(
            MARKET_VALUE_URL,
            params={"page": page, "pageSize": PAGE_SIZE},
            headers=HEADERS,
            timeout=10,
        )
        res.raise_for_status()
        data = res.json()
        page_stocks = data.get("stocks", [])
        if not page_stocks:
            break

        stocks.extend(s for s in page_stocks if s.get("stockEndType") == "stock")
        if progress_callback:
            progress_callback(f"일반 주식 수집 중... ({min(len(stocks), TARGET_COUNT)}/{TARGET_COUNT})")
        page += 1

    stocks = stocks[:TARGET_COUNT]

    results = []
    for rank, s in enumerate(stocks, start=1):
        sign = SIGN_BY_TEXT.get(s.get("compareToPreviousPrice", {}).get("text", ""), "")
        code = s.get("itemCode", "")
        results.append({
            "rank": rank,
            "code": code,
            "name": s.get("stockName", ""),
            "price": s.get("closePrice", ""),
            "change": f"{sign}{s.get('compareToPreviousClosePrice', '')}",
            "change_rate": f"{sign}{s.get('fluctuationsRatio', '')}%",
            "direction": s.get("compareToPreviousPrice", {}).get("text", ""),
            "volume": s.get("accumulatedTradingVolume", ""),
            "market_cap": s.get("marketValue", ""),
            "link": DETAIL_URL_TEMPLATE.format(code=code),
        })
    return results


class FetchWorker(QThread):
    """네트워크 요청을 백그라운드 스레드에서 수행해 UI 멈춤을 막는다."""

    progress = pyqtSignal(str)
    finished_all = pyqtSignal(list)
    error = pyqtSignal(str)

    def run(self):
        try:
            results = fetch_kospi200(progress_callback=self.progress.emit)
        except requests.RequestException as e:
            self.error.emit(f"시세 데이터를 가져오지 못했습니다: {e}")
            return
        except ValueError as e:
            self.error.emit(f"응답 데이터를 해석하지 못했습니다: {e}")
            return

        self.finished_all.emit(results)


class MainWindow(QMainWindow):
    COLUMNS = ["순위", "종목코드", "종목명", "현재가", "전일대비", "등락률", "거래량", "시가총액(억원)"]

    def __init__(self):
        super().__init__()
        self.setWindowTitle("코스피 200 시세")
        self.resize(1000, 640)

        self.results = []
        self.worker = None

        self._build_ui()
        self.start_fetch()

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        layout = QVBoxLayout(central)

        top_row = QHBoxLayout()
        self.refresh_btn = QPushButton("새로고침")
        self.refresh_btn.clicked.connect(self.start_fetch)
        top_row.addWidget(self.refresh_btn)

        top_row.addWidget(QLabel("종목명/코드 검색"))
        self.filter_edit = QLineEdit()
        self.filter_edit.textChanged.connect(self.apply_filter)
        top_row.addWidget(self.filter_edit, stretch=1)

        self.save_csv_btn = QPushButton("CSV로 저장")
        self.save_csv_btn.setEnabled(False)
        self.save_csv_btn.clicked.connect(self.save_csv)
        top_row.addWidget(self.save_csv_btn)

        self.save_excel_btn = QPushButton("엑셀로 저장")
        self.save_excel_btn.setEnabled(False)
        self.save_excel_btn.clicked.connect(self.save_excel)
        top_row.addWidget(self.save_excel_btn)

        layout.addLayout(top_row)

        self.table = QTableWidget(0, len(self.COLUMNS))
        self.table.setHorizontalHeaderLabels(self.COLUMNS)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        layout.addWidget(self.table, stretch=1)

        self.status_label = QLabel("대기 중")
        layout.addWidget(self.status_label)

    def start_fetch(self):
        if self.worker is not None and self.worker.isRunning():
            return

        self.table.setRowCount(0)
        self.results = []
        self.save_csv_btn.setEnabled(False)
        self.save_excel_btn.setEnabled(False)
        self.refresh_btn.setEnabled(False)
        self.refresh_btn.setText("수집 중...")

        self.worker = FetchWorker()
        self.worker.progress.connect(self.on_progress)
        self.worker.finished_all.connect(self.on_finished)
        self.worker.error.connect(self.on_error)
        self.worker.start()

    def on_progress(self, text):
        self.status_label.setText(text)

    def on_finished(self, results):
        self.results = results
        self.table.setRowCount(len(results))
        for row, item in enumerate(results):
            self._set_row(row, item)

        self.status_label.setText(f"완료: {len(results)}개 종목 (코스피 시가총액 상위 기준)")
        self.refresh_btn.setEnabled(True)
        self.refresh_btn.setText("새로고침")
        self.save_csv_btn.setEnabled(bool(results))
        self.save_excel_btn.setEnabled(bool(results))
        self.apply_filter(self.filter_edit.text())

    def on_error(self, message):
        QMessageBox.critical(self, "오류", message)
        self.status_label.setText("오류로 중단됨")
        self.refresh_btn.setEnabled(True)
        self.refresh_btn.setText("새로고침")

    def _set_row(self, row, item):
        values = [
            str(item["rank"]),
            item["code"],
            item["name"],
            item["price"],
            item["change"],
            item["change_rate"],
            item["volume"],
            item["market_cap"],
        ]
        color = QColor("#d60000") if item["direction"] == "상승" else (
            QColor("#0051c7") if item["direction"] == "하락" else None
        )

        for col, value in enumerate(values):
            cell = QTableWidgetItem(value)
            if col in (0, 3, 4, 5, 6, 7):
                cell.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            if color is not None and col in (4, 5):
                cell.setForeground(color)
            self.table.setItem(row, col, cell)

    def apply_filter(self, text):
        text = text.strip().lower()
        for row in range(self.table.rowCount()):
            if not text:
                self.table.setRowHidden(row, False)
                continue
            name = self.table.item(row, 2).text().lower()
            code = self.table.item(row, 1).text().lower()
            self.table.setRowHidden(row, text not in name and text not in code)

    def save_csv(self):
        if not self.results:
            return
        path, _ = QFileDialog.getSaveFileName(self, "CSV로 저장", "kospi200.csv", "CSV Files (*.csv)")
        if not path:
            return

        fieldnames = ["rank", "code", "name", "price", "change", "change_rate", "volume", "market_cap", "link"]
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(self.results)

        QMessageBox.information(self, "저장 완료", f"{path}\n에 저장되었습니다.")

    def save_excel(self):
        if not self.results:
            return
        path, _ = QFileDialog.getSaveFileName(self, "엑셀로 저장", "kospi200.xlsx", "Excel Files (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        headers = self.COLUMNS + ["링크"]
        keys = ["rank", "code", "name", "price", "change", "change_rate", "volume", "market_cap", "link"]

        wb = Workbook()
        ws = wb.active
        ws.title = "코스피200"
        ws.append(headers)

        for item in self.results:
            ws.append([item[key] for key in keys])

        widths = [6, 10, 16, 12, 12, 10, 14, 16, 42]
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

        wb.save(path)
        QMessageBox.information(self, "저장 완료", f"{path}\n에 저장되었습니다.")

    def closeEvent(self, event):
        if self.worker is not None and self.worker.isRunning():
            self.worker.wait(3000)
        event.accept()


def main():
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()

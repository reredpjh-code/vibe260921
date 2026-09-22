"""
PyQt6 GUI 기반 마켓컬리 상품 후기 감성분석기

상품 URL(또는 상품번호)을 입력하면 마켓컬리 후기 API에서 후기를 약 300개
수집하고, Pandas로 정리한 뒤 간단한 한국어 감성 사전으로 긍정/부정/중립을
분류해 matplotlib 파이차트로 보여준다. 결과는 CSV 또는 엑셀(.xlsx)로 저장할
수 있다.

마켓컬리 후기 API(product-review/v4)는 별점 필드를 제공하지 않으므로, 후기
본문에 담긴 감성 단어 개수를 세어 긍정/부정/중립을 판정하는 규칙 기반
방식을 쓴다. 형태소 분석기 없이 문자열 포함 여부만으로 판정하기 때문에
부정어("안", "못")로 뒤집힌 문장은 오판할 수 있다는 한계가 있지만, 별도
모델 설치 없이 바로 동작한다.
"""
import csv
import html
import os
import re
import sys
import time

os.environ.setdefault("QT_API", "pyqt6")

import matplotlib
import matplotlib.font_manager as fm
import pandas as pd
import requests
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg
from matplotlib.figure import Figure
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QComboBox,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept": "application/json",
}

REVIEW_API_URL = "https://api.kurly.com/product-review/v4/contents-products/{goods_no}/reviews"
PER_PAGE = 10

DEFAULT_URL = "https://www.kurly.com/goods/1000969096?collectionCode=sale231107"

# matplotlib 기본 폰트(DejaVu Sans)는 한글 글리프가 없어 텍스트가 네모(□)로
# 깨진다. OS별로 흔히 설치돼 있는 한글 폰트를 찾아서 지정한다.
_KOREAN_FONT_CANDIDATES = [
    "Apple SD Gothic Neo", "AppleGothic",  # macOS
    "Malgun Gothic",  # Windows
    "NanumGothic", "NanumBarunGothic", "Noto Sans CJK KR", "Noto Sans KR",  # Linux
]
_installed_fonts = {f.name for f in fm.fontManager.ttflist}
for _font in _KOREAN_FONT_CANDIDATES:
    if _font in _installed_fonts:
        matplotlib.rcParams["font.family"] = _font
        break
matplotlib.rcParams["axes.unicode_minus"] = False

# 상태 팔레트(긍정/부정/중립)는 데이터의 "의미"를 나타내는 색이므로 카테고리
# 색상이 아닌 상태(status) 색상을 쓴다: 긍정=good, 부정=critical, 중립=회색.
SENTIMENT_COLORS = {
    "긍정": "#0ca30c",
    "부정": "#d03b3b",
    "중립": "#898781",
}

POSITIVE_WORDS = [
    "맛있", "맛나", "존맛", "꿀맛", "부드럽", "촉촉", "신선", "싱싱",
    "최고", "훌륭", "만족", "좋아요", "좋았어요", "좋네요", "좋습니다",
    "추천", "강추", "대박", "깔끔", "든든", "알차", "저렴", "가성비",
    "빠른 배송", "빠르게 와", "친절", "완벽", "쫄깃", "고소", "담백",
    "육즙", "야들", "부들", "최애", "감사", "행복", "편리", "편해요",
    "믿고 먹", "믿고 구매", "다시 구매", "또 구매", "또 살", "재구매",
    "진짜 맛있", "정말 좋", "너무 좋", "항상 만족", "꾸준히 구매", "단골",
    "실망하지 않", "성공적",
]

NEGATIVE_WORDS = [
    "별로", "아쉬", "실망", "최악", "비려", "비린내", "잡내", "냄새나",
    "질겨", "질기", "딱딱", "퍽퍽", "짜요", "싱거", "느끼", "비싸",
    "부족", "작아요", "적어요", "상했", "변질", "곰팡이", "썩",
    "파손", "터졌", "터져", "늦게 와", "늦었", "배송 지연", "환불",
    "교환", "불친절", "불만", "화나", "짜증", "다시는", "구매 후회",
    "후회", "그닥", "그저 그", "실패", "안 좋", "맛없", "밍밍", "물컹",
    "핏물", "비위생", "불량", "찢어", "눅눅",
]


def extract_goods_no(text):
    """URL 또는 숫자 문자열에서 마켓컬리 상품번호를 추출한다."""
    text = text.strip()
    if text.isdigit():
        return text
    m = re.search(r"/goods/(\d+)", text)
    if m:
        return m.group(1)
    raise ValueError("상품 URL 또는 상품번호를 확인해주세요. 예: https://www.kurly.com/goods/1000969096")


def fetch_reviews(goods_no, target_count, progress_cb=None, should_stop=None):
    """커서 기반 페이지네이션(after)으로 후기를 target_count개까지 수집한다."""
    url = REVIEW_API_URL.format(goods_no=goods_no)
    reviews = []
    after = None

    while len(reviews) < target_count:
        if should_stop and should_stop():
            break

        params = {"per_page": PER_PAGE, "sort": "recent"}
        if after:
            params["after"] = after

        res = requests.get(url, params=params, headers=HEADERS, timeout=10)
        res.raise_for_status()
        payload = res.json()
        if not payload.get("success"):
            raise RuntimeError(payload.get("message") or "후기를 가져오지 못했습니다.")

        batch = payload["data"]["reviews"]
        if not batch:
            break

        reviews.extend(batch)
        if progress_cb:
            progress_cb(len(reviews))

        after = (payload["data"].get("nextCursor") or {}).get("after")
        if not after:
            break

        time.sleep(0.25)

    return reviews[:target_count]


def analyze_sentiment(text):
    """긍정/부정 단어 개수를 세어 다수인 쪽으로 분류한다. 동률이면 중립."""
    pos = sum(1 for w in POSITIVE_WORDS if w in text)
    neg = sum(1 for w in NEGATIVE_WORDS if w in text)
    if pos > neg:
        label = "긍정"
    elif neg > pos:
        label = "부정"
    else:
        label = "중립"
    return label, pos, neg


def review_to_row(review):
    text = review.get("contents", "") or ""
    label, pos, neg = analyze_sentiment(text)
    return {
        "번호": review.get("no"),
        "작성자": review.get("ownerName", ""),
        "상품명": review.get("contentsProductName") or review.get("dealProductName") or "",
        "내용": text,
        "좋아요": review.get("likeCount", 0),
        "작성일": (review.get("registeredAt") or "")[:10],
        "감정": label,
        "긍정단어수": pos,
        "부정단어수": neg,
    }


def build_representative_html(df, top_n=5):
    """감정별로 좋아요(도움돼요) 수가 높은 후기를 뽑아 대표 의견으로 인용한다."""
    if df.empty:
        return "<i>후기를 수집하면 대표 의견이 표시됩니다.</i>"

    sections = []
    for label in ["긍정", "부정", "중립"]:
        sub = df[df["감정"] == label]
        if sub.empty:
            continue
        score_col = "긍정단어수" if label == "긍정" else "부정단어수"
        top = sub.sort_values(["좋아요", score_col], ascending=False).head(top_n)

        color = SENTIMENT_COLORS[label]
        sections.append(
            f'<h3 style="color:{color}; margin-bottom:4px;">{label} 대표 의견 ({len(sub)}개 중)</h3>'
        )
        for _, row in top.iterrows():
            text = html.escape(row["내용"]).replace("\n", "<br>")
            sections.append(
                '<blockquote style="margin:0 0 10px 0; padding:6px 10px; '
                f'border-left:3px solid {color}; color:#0b0b0b;">'
                f"“{text}”<br>"
                f'<span style="color:#898781; font-size:11px;">'
                f"— {html.escape(row['작성자'])} · {row['작성일']} · 도움돼요 {row['좋아요']}"
                "</span></blockquote>"
            )
    return "".join(sections)


class CrawlWorker(QThread):
    """네트워크 요청과 감성분석을 백그라운드 스레드에서 수행해 UI 멈춤을 막는다."""

    progress = pyqtSignal(str)
    item_crawled = pyqtSignal(dict)
    finished_all = pyqtSignal(list)
    error = pyqtSignal(str)

    def __init__(self, goods_no, target_count):
        super().__init__()
        self.goods_no = goods_no
        self.target_count = target_count
        self._stop = False

    def stop(self):
        self._stop = True

    def run(self):
        try:
            reviews = fetch_reviews(
                self.goods_no,
                self.target_count,
                progress_cb=lambda n: self.progress.emit(f"후기 {n}개 수집 중..."),
                should_stop=lambda: self._stop,
            )
        except (requests.RequestException, RuntimeError) as e:
            self.error.emit(f"후기를 가져오지 못했습니다: {e}")
            return

        if not reviews:
            self.error.emit("수집된 후기가 없습니다. 상품번호를 확인해주세요.")
            return

        rows = []
        for review in reviews:
            if self._stop:
                break
            row = review_to_row(review)
            rows.append(row)
            self.item_crawled.emit(row)

        self.finished_all.emit(rows)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("마켓컬리 후기 감성분석기")
        self.resize(1200, 720)

        self.df = pd.DataFrame()
        self.worker = None

        self._build_ui()

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        layout = QVBoxLayout(central)

        input_row = QHBoxLayout()
        input_row.addWidget(QLabel("상품 URL / 상품번호"))
        self.url_edit = QLineEdit(DEFAULT_URL)
        self.url_edit.returnPressed.connect(self.start_crawl)
        input_row.addWidget(self.url_edit, stretch=1)

        input_row.addWidget(QLabel("수집할 후기 수"))
        self.count_spin = QSpinBox()
        self.count_spin.setRange(10, 2000)
        self.count_spin.setSingleStep(10)
        self.count_spin.setValue(300)
        input_row.addWidget(self.count_spin)

        self.start_btn = QPushButton("수집 및 분석 시작")
        self.start_btn.clicked.connect(self.start_crawl)
        input_row.addWidget(self.start_btn)

        self.save_csv_btn = QPushButton("CSV로 저장")
        self.save_csv_btn.setEnabled(False)
        self.save_csv_btn.clicked.connect(self.save_csv)
        input_row.addWidget(self.save_csv_btn)

        self.save_excel_btn = QPushButton("엑셀로 저장")
        self.save_excel_btn.setEnabled(False)
        self.save_excel_btn.clicked.connect(self.save_excel)
        input_row.addWidget(self.save_excel_btn)

        layout.addLayout(input_row)

        splitter = QSplitter()

        # 왼쪽: 후기 목록(감정별 필터) + 선택한 후기 원문 미리보기
        left = QWidget()
        left_layout = QVBoxLayout(left)
        left_layout.setContentsMargins(0, 0, 0, 0)

        filter_row = QHBoxLayout()
        filter_row.addWidget(QLabel("감정 필터"))
        self.filter_combo = QComboBox()
        self.filter_combo.addItems(["전체", "긍정", "부정", "중립"])
        self.filter_combo.currentTextChanged.connect(self.on_filter_changed)
        filter_row.addWidget(self.filter_combo)
        filter_row.addStretch(1)
        left_layout.addLayout(filter_row)

        self.table = QTableWidget(0, 6)
        self.table.setHorizontalHeaderLabels(["감정", "작성자", "상품명", "내용", "좋아요", "작성일"])
        self.table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.itemSelectionChanged.connect(self.show_selected_detail)
        left_layout.addWidget(self.table, stretch=3)

        left_layout.addWidget(QLabel("선택한 후기 원문"))
        self.detail_view = QTextEdit()
        self.detail_view.setReadOnly(True)
        left_layout.addWidget(self.detail_view, stretch=1)

        splitter.addWidget(left)

        # 오른쪽: 파이차트 + 대표 의견(원문 인용)
        right = QWidget()
        right_layout = QVBoxLayout(right)
        right_layout.setContentsMargins(0, 0, 0, 0)

        self.figure = Figure(figsize=(4.5, 4.5), facecolor="#fcfcfb")
        self.canvas = FigureCanvasQTAgg(self.figure)
        self._draw_empty_chart()
        right_layout.addWidget(self.canvas, stretch=2)

        right_layout.addWidget(QLabel("대표 의견 (도움돼요 순)"))
        self.representative_view = QTextEdit()
        self.representative_view.setReadOnly(True)
        self.representative_view.setHtml(build_representative_html(pd.DataFrame()))
        right_layout.addWidget(self.representative_view, stretch=2)

        splitter.addWidget(right)

        splitter.setStretchFactor(0, 3)
        splitter.setStretchFactor(1, 2)
        layout.addWidget(splitter, stretch=1)

        self.status_label = QLabel("대기 중")
        layout.addWidget(self.status_label)

    def _draw_empty_chart(self):
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        ax.text(0.5, 0.5, "후기를 수집하면\n감성 분석 결과가 표시됩니다",
                 ha="center", va="center", color="#52514e", fontsize=11)
        ax.axis("off")
        self.canvas.draw()

    def start_crawl(self):
        if self.worker is not None and self.worker.isRunning():
            return

        try:
            goods_no = extract_goods_no(self.url_edit.text())
        except ValueError as e:
            QMessageBox.warning(self, "입력 오류", str(e))
            return

        self.table.setRowCount(0)
        self.detail_view.clear()
        self.df = pd.DataFrame()
        self.filter_combo.blockSignals(True)
        self.filter_combo.setCurrentIndex(0)
        self.filter_combo.blockSignals(False)
        self.representative_view.setHtml(build_representative_html(self.df))
        self.save_csv_btn.setEnabled(False)
        self.save_excel_btn.setEnabled(False)
        self.start_btn.setEnabled(False)
        self.start_btn.setText("수집 중...")
        self._draw_empty_chart()

        self.worker = CrawlWorker(goods_no, self.count_spin.value())
        self.worker.progress.connect(self.on_progress)
        self.worker.item_crawled.connect(self.on_item_crawled)
        self.worker.finished_all.connect(self.on_finished)
        self.worker.error.connect(self.on_error)
        self.worker.start()

    def on_progress(self, text):
        self.status_label.setText(text)

    def on_item_crawled(self, row):
        self._append_table_row(row)

    def _append_table_row(self, row):
        i = self.table.rowCount()
        self.table.insertRow(i)
        self.table.setItem(i, 0, QTableWidgetItem(row["감정"]))
        self.table.setItem(i, 1, QTableWidgetItem(row["작성자"]))
        self.table.setItem(i, 2, QTableWidgetItem(row["상품명"]))
        self.table.setItem(i, 3, QTableWidgetItem(row["내용"]))
        self.table.setItem(i, 4, QTableWidgetItem(str(row["좋아요"])))
        self.table.setItem(i, 5, QTableWidgetItem(row["작성일"]))
        # 상세 미리보기에서 원본 후기 전체를 그대로 보여주기 위해 행 데이터를 보관한다.
        self.table.item(i, 0).setData(Qt.ItemDataRole.UserRole, row)

    def populate_table(self, rows):
        self.table.setRowCount(0)
        for row in rows:
            self._append_table_row(row)
        self.detail_view.clear()

    def on_filter_changed(self, label):
        if self.df.empty:
            return
        if label == "전체":
            rows = self.df.to_dict("records")
        else:
            rows = self.df[self.df["감정"] == label].to_dict("records")
        self.populate_table(rows)

    def show_selected_detail(self):
        selected = self.table.selectedItems()
        if not selected:
            return
        row_index = selected[0].row()
        row = self.table.item(row_index, 0).data(Qt.ItemDataRole.UserRole)
        if not row:
            return
        self.detail_view.setPlainText(
            f"[{row['감정']}] {row['상품명']}\n"
            f"작성자: {row['작성자']}  |  작성일: {row['작성일']}  |  도움돼요: {row['좋아요']}\n"
            f"{'-' * 40}\n"
            f"{row['내용']}"
        )

    def on_finished(self, rows):
        self.df = pd.DataFrame(rows)
        self.status_label.setText(f"완료: 후기 {len(self.df)}개 수집 및 분석됨")
        self.start_btn.setEnabled(True)
        self.start_btn.setText("수집 및 분석 시작")
        self.save_csv_btn.setEnabled(not self.df.empty)
        self.save_excel_btn.setEnabled(not self.df.empty)
        self.populate_table(rows)
        self.representative_view.setHtml(build_representative_html(self.df))
        self.draw_pie_chart()

    def on_error(self, message):
        QMessageBox.critical(self, "오류", message)
        self.status_label.setText("오류로 중단됨")
        self.start_btn.setEnabled(True)
        self.start_btn.setText("수집 및 분석 시작")

    def draw_pie_chart(self):
        if self.df.empty:
            self._draw_empty_chart()
            return

        counts = self.df["감정"].value_counts()
        order = [k for k in ["긍정", "부정", "중립"] if k in counts.index]
        values = [counts[k] for k in order]
        colors = [SENTIMENT_COLORS[k] for k in order]
        total = sum(values)

        self.figure.clear()
        ax = self.figure.add_subplot(111)

        def autopct(pct):
            n = round(pct * total / 100)
            return f"{pct:.1f}%\n({n}개)"

        wedges, _texts, _autotexts = ax.pie(
            values,
            colors=colors,
            autopct=autopct,
            startangle=90,
            counterclock=False,
            wedgeprops={"edgecolor": "#fcfcfb", "linewidth": 2},
            textprops={"color": "#0b0b0b", "fontsize": 10},
        )
        ax.set_title(f"후기 감성 분석 결과 (총 {total}개)", color="#0b0b0b", fontsize=12)
        ax.legend(
            wedges,
            [f"{k} ({counts[k]}개)" for k in order],
            loc="upper center",
            bbox_to_anchor=(0.5, -0.02),
            ncol=len(order),
            frameon=False,
        )
        ax.axis("equal")
        self.figure.tight_layout()
        self.canvas.draw()

    def save_csv(self):
        if self.df.empty:
            return
        path, _ = QFileDialog.getSaveFileName(self, "CSV로 저장", "kurly_reviews.csv", "CSV Files (*.csv)")
        if not path:
            return
        self.df.to_csv(path, index=False, encoding="utf-8-sig", quoting=csv.QUOTE_MINIMAL)
        QMessageBox.information(self, "저장 완료", f"{path}\n에 저장되었습니다.")

    def save_excel(self):
        if self.df.empty:
            return
        path, _ = QFileDialog.getSaveFileName(self, "엑셀로 저장", "kurly_reviews.xlsx", "Excel Files (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "컬리후기"
        ws.append(list(self.df.columns))
        for row in self.df.itertuples(index=False):
            ws.append(list(row))

        widths = [10, 10, 30, 60, 8, 12, 8, 10, 10]
        for i, width in enumerate(widths[: len(self.df.columns)], start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

        wb.save(path)
        QMessageBox.information(self, "저장 완료", f"{path}\n에 저장되었습니다.")

    def closeEvent(self, event):
        if self.worker is not None and self.worker.isRunning():
            self.worker.stop()
            self.worker.wait(3000)
        event.accept()


def main():
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
